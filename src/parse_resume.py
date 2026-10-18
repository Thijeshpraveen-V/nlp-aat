"""PDF/DOCX/TXT -> text -> entities -> structured JSON. Shared by the API and the demo.

    python src/parse_resume.py data/sample_resumes/sample_fictional.pdf --model finetuned
"""
import argparse
import io
import json
import os
import re
import time
from functools import lru_cache
from pathlib import Path

import torch
from transformers import AutoModelForTokenClassification, AutoTokenizer

from common import BASE_MODEL, MODELS, bio_to_spans, word_tokenize
from data_prep import clean_text
from evaluate import repair_bio
from inference import predict_word_probs, probs_to_tags
from pretrained_eval import map_tags
from rules import apply_rules, find_phones

FINETUNED_DIR = Path(os.environ.get("RESUME_NER_MODEL", MODELS / "main"))
MODEL_NAMES = ("pretrained", "finetuned")


def extract_text(data: bytes, filename: str) -> str:
    return _extract(data, filename).replace("\r\n", "\n").replace("\r", "\n")


def _extract(data: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [p.extract_text() or "" for p in pdf.pages]
        text = "\n".join(pages)
        if not text.strip():
            raise ValueError("No extractable text in PDF (scanned/image-only PDFs need OCR).")
        return text
    if ext == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext in (".txt", ".text", ""):
        return data.decode("utf-8", errors="replace")
    raise ValueError(f"Unsupported file type '{ext}'. Use PDF, DOCX or TXT.")


@lru_cache(maxsize=2)
def load_model(name: str):
    if name == "pretrained":
        path = BASE_MODEL
    elif name == "finetuned":
        if not FINETUNED_DIR.exists():
            raise FileNotFoundError(f"Fine-tuned model not found at {FINETUNED_DIR}. Run src/finetune.py --save_model.")
        path = str(FINETUNED_DIR)
    else:
        raise ValueError(f"model must be one of {MODEL_NAMES}")
    tok = AutoTokenizer.from_pretrained(path)
    model = AutoModelForTokenClassification.from_pretrained(path).eval()
    return tok, model


def split_items(texts):
    """Split SKILLS blocks into individual items (commas, bullets, newlines)."""
    items = []
    for t in texts:
        for part in re.split(r"[,\n;•|]+", t):
            part = re.sub(r"\(.*?\)", "", part).strip(" -:\t")
            if part and part.lower() not in {i.lower() for i in items}:
                items.append(part)
    return items


def structure(entities, text):
    by = {}
    for e in entities:
        by.setdefault(e["label"], [])
        if e["text"] not in by[e["label"]]:
            by[e["label"]].append(e["text"])
    first = lambda k: by.get(k, [None])[0]
    return {
        "name": first("NAME"),
        "email": first("EMAIL"),
        "phone": (find_phones(text) or [None])[0],
        "location": by.get("LOCATION", []),
        "skills": split_items(by.get("SKILLS", [])),
        "education": {"degrees": by.get("DEGREE", []), "colleges": by.get("COLLEGE_NAME", []),
                      "graduation_years": by.get("GRADUATION_YEAR", [])},
        "companies": by.get("COMPANIES_WORKED_AT", []),
        "designations": by.get("DESIGNATION", []),
        "years_of_experience": by.get("YEARS_OF_EXPERIENCE", []),
    }


def parse_text(text: str, model_name: str = "finetuned", rules: bool = True):
    tok, model = load_model(model_name)  # cached; load time excluded from latency
    t0 = time.perf_counter()
    text = clean_text(text)
    words, starts, ends = word_tokenize(text)
    if not words:
        return {"model": model_name, "text": text, "entities": [], "latency_ms": 0.0, **structure([], text)}
    with torch.inference_mode():
        probs = predict_word_probs(model, tok, words, device="cpu")
    tags = probs_to_tags(probs, model.config.id2label)
    tags = map_tags(tags) if model_name == "pretrained" else repair_bio(tags)
    if rules and model_name == "finetuned":
        tags = repair_bio(apply_rules(text, starts, ends, tags))
    entities = [{"text": text[s:e], "label": t, "start": s, "end": e} for s, e, t in bio_to_spans(tags, starts, ends)]
    out = structure(entities, text)
    return {"model": model_name, "rules": rules and model_name == "finetuned", "text": text,
            "n_words": len(words), "entities": entities, "latency_ms": (time.perf_counter() - t0) * 1000, **out}


def parse_file(data: bytes, filename: str, model_name="finetuned", rules=True):
    return parse_text(extract_text(data, filename), model_name, rules)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--model", default="finetuned", choices=MODEL_NAMES)
    ap.add_argument("--no_rules", action="store_true")
    a = ap.parse_args()
    r = parse_file(Path(a.path).read_bytes(), a.path, a.model, not a.no_rules)
    r.pop("text")
    print(json.dumps(r, indent=2, ensure_ascii=False))
